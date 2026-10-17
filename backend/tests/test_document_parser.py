"""Comprehensive Tests for Document Ingestion Layer (POST /api/v1/documents/parse).

Tests:
- PDF text extraction via PyMuPDF
- DOCX text extraction via python-docx
- TXT plain text extraction
- Empty file rejection (400)
- Unsupported file format rejection (400)
- File size limit enforcement (413)
- Corrupted PDF handling (422)
- Corrupted DOCX handling (422)
- Password-protected document handling (422)
- Sensitive data sanitization (Aadhaar / card masking)
- End-to-end FastAPI endpoint validation
"""

import io
import pytest
import fitz  # PyMuPDF
import docx
from fastapi.testclient import TestClient

from app.main import app
from app.services.document_parser import (
    parse_document,
    DocumentParseResponse,
    EmptyFileError,
    FileSizeExceededError,
    UnsupportedFormatError,
    EncryptedDocumentError,
    CorruptedDocumentError,
    MAX_FILE_SIZE_BYTES,
)


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


def create_sample_pdf_bytes(text: str = "Pawan Kumar\nSenior Backend Engineer\nPython, FastAPI, Docker, SQL") -> bytes:
    """Create in-memory valid PDF bytes using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes(
    paragraphs: list[str] = None,
    table_data: list[list[str]] = None
) -> bytes:
    """Create in-memory valid DOCX bytes using python-docx."""
    if paragraphs is None:
        paragraphs = [
            "Aarav Sharma",
            "Full Stack Developer with 4 years experience",
            "Skills: Python, TypeScript, React, PostgreSQL",
        ]
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)

    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, cell in enumerate(row):
                table.rows[r_idx].cells[c_idx].text = cell

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def create_encrypted_pdf_bytes(password: str = "SecretPass123!") -> bytes:
    """Create in-memory password-protected PDF bytes."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Confidential Candidate Resume Details")
    buffer = io.BytesIO()
    doc.save(
        buffer,
        encryption=fitz.PDF_ENCRYPT_AES_256,
        user_pw=password,
        permissions=fitz.PDF_PERM_ACCESSIBILITY,
    )
    doc.close()
    return buffer.getvalue()


# ==============================================================================
# 1. UNIT TESTS: DOCUMENT PARSER SERVICE
# ==============================================================================

def test_parse_valid_pdf():
    pdf_bytes = create_sample_pdf_bytes("Vikram Verma\nMachine Learning Engineer\nPyTorch, Python, Docker")
    res = parse_document(pdf_bytes, "resume.pdf", "application/pdf")

    assert isinstance(res, DocumentParseResponse)
    assert res.document_type == "pdf"
    assert "Vikram Verma" in res.text
    assert "PyTorch" in res.text
    assert res.page_count == 1
    assert res.metadata["filename"] == "resume.pdf"
    assert res.metadata["word_count"] >= 5
    assert res.metadata["file_size_bytes"] == len(pdf_bytes)


def test_parse_valid_docx():
    docx_bytes = create_sample_docx_bytes(
        paragraphs=["Priya Patel", "Data Scientist", "Python, Statistics, SQL, Pandas"],
        table_data=[["Company", "Role", "Years"], ["TechCorp", "Lead Data Scientist", "3"]]
    )
    res = parse_document(docx_bytes, "priya_resume.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    assert res.document_type == "docx"
    assert "Priya Patel" in res.text
    assert "Lead Data Scientist" in res.text
    assert res.page_count >= 1
    assert res.metadata["filename"] == "priya_resume.docx"
    assert res.metadata["word_count"] >= 6


def test_parse_valid_txt():
    content = "Rohan Gupta\nFrontend Developer\nReact, Next.js, TypeScript"
    txt_bytes = content.encode("utf-8")
    res = parse_document(txt_bytes, "profile.txt", "text/plain")

    assert res.document_type == "txt"
    assert "Rohan Gupta" in res.text
    assert "TypeScript" in res.text
    assert res.page_count == 1
    assert res.metadata["word_count"] >= 5


def test_empty_file_raises_error():
    with pytest.raises(EmptyFileError) as exc_info:
        parse_document(b"", "empty.pdf")
    assert "empty" in str(exc_info.value).lower()


def test_unsupported_format_raises_error():
    with pytest.raises(UnsupportedFormatError) as exc_info:
        parse_document(b"fake image bytes", "photo.png")
    assert "unsupported" in str(exc_info.value).lower()
    assert ".png" in str(exc_info.value)


def test_file_size_limit_exceeded():
    large_bytes = b"0" * (MAX_FILE_SIZE_BYTES + 1024)
    with pytest.raises(FileSizeExceededError) as exc_info:
        parse_document(large_bytes, "huge_resume.pdf")
    assert "exceeds" in str(exc_info.value).lower()


def test_corrupted_pdf_raises_error():
    corrupted_bytes = b"%PDF-1.4 ... totally garbage corrupted bytes not valid PDF"
    with pytest.raises(CorruptedDocumentError) as exc_info:
        parse_document(corrupted_bytes, "broken.pdf")
    assert "corrupted" in str(exc_info.value).lower()


def test_corrupted_docx_raises_error():
    corrupted_bytes = b"PK\x03\x04 corrupt zip archive data"
    with pytest.raises(CorruptedDocumentError) as exc_info:
        parse_document(corrupted_bytes, "broken.docx")
    assert "corrupted" in str(exc_info.value).lower()


def test_encrypted_pdf_raises_error():
    encrypted_bytes = create_encrypted_pdf_bytes("Secret123")
    with pytest.raises(EncryptedDocumentError) as exc_info:
        parse_document(encrypted_bytes, "locked.pdf")
    assert "password-protected" in str(exc_info.value).lower()


def test_sensitive_information_masked():
    text_with_pii = (
        "Candidate Name: Rajesh Kumar\n"
        "Aadhaar Number: 4521 8902 3412\n"
        "Credit Card: 4111-2222-3333-4444\n"
        "Technical Skills: Python, SQL"
    )
    txt_bytes = text_with_pii.encode("utf-8")
    res = parse_document(txt_bytes, "pii.txt")

    assert "[REDACTED-NATIONAL-ID]" in res.text
    assert "4521 8902 3412" not in res.text
    assert "[REDACTED-PAYMENT-CARD]" in res.text
    assert "4111-2222-3333-4444" not in res.text
    assert "Rajesh Kumar" in res.text
    assert "Python, SQL" in res.text


# ==============================================================================
# 2. ENDPOINT INTEGRATION TESTS (POST /api/v1/documents/parse)
# ==============================================================================

def test_endpoint_parse_pdf_success(client: TestClient):
    pdf_bytes = create_sample_pdf_bytes("Ananya Iyer\nDevOps Engineer\nAWS, Terraform, Kubernetes, Linux")
    files = {"file": ("ananya_resume.pdf", pdf_bytes, "application/pdf")}

    response = client.post("/api/v1/documents/parse", files=files)
    assert response.status_code == 200

    data = response.json()
    assert data["document_type"] == "pdf"
    assert "Ananya Iyer" in data["text"]
    assert "Terraform" in data["text"]
    assert data["page_count"] == 1
    assert data["metadata"]["filename"] == "ananya_resume.pdf"
    assert data["metadata"]["word_count"] >= 6


def test_endpoint_parse_docx_success(client: TestClient):
    docx_bytes = create_sample_docx_bytes(
        paragraphs=["Siddharth Rao", "Applied ML Specialist", "PyTorch, HuggingFace, Transformers"]
    )
    files = {
        "file": (
            "siddharth.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }

    response = client.post("/api/v1/documents/parse", files=files)
    assert response.status_code == 200

    data = response.json()
    assert data["document_type"] == "docx"
    assert "Siddharth Rao" in data["text"]
    assert "Transformers" in data["text"]
    assert data["metadata"]["filename"] == "siddharth.docx"


def test_endpoint_parse_txt_success(client: TestClient):
    txt_bytes = b"Neha Kulkarni\nBusiness Intelligence\nPower BI, SQL, Excel"
    files = {"file": ("neha.txt", txt_bytes, "text/plain")}

    response = client.post("/api/v1/documents/parse", files=files)
    assert response.status_code == 200

    data = response.json()
    assert data["document_type"] == "txt"
    assert "Neha Kulkarni" in data["text"]


def test_endpoint_empty_file_rejected(client: TestClient):
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    response = client.post("/api/v1/documents/parse", files=files)
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_endpoint_unsupported_format_rejected(client: TestClient):
    files = {"file": ("malicious.exe", b"fake binary executable content", "application/octet-stream")}
    response = client.post("/api/v1/documents/parse", files=files)
    assert response.status_code == 400
    assert "unsupported" in response.json()["detail"].lower()


def test_endpoint_corrupted_pdf_rejected(client: TestClient):
    files = {"file": ("corrupted.pdf", b"Not a real PDF stream", "application/pdf")}
    response = client.post("/api/v1/documents/parse", files=files)
    assert response.status_code == 422
    assert "corrupted" in response.json()["detail"].lower()


def test_endpoint_encrypted_pdf_rejected(client: TestClient):
    locked_pdf = create_encrypted_pdf_bytes("SecretKey")
    files = {"file": ("locked.pdf", locked_pdf, "application/pdf")}
    response = client.post("/api/v1/documents/parse", files=files)
    assert response.status_code == 422
    assert "password-protected" in response.json()["detail"].lower()
